from django.shortcuts import render, redirect
from .forms import MessageForm
from .models import Message
from django.core.files.storage import FileSystemStorage
from .models import TestResult
import openpyxl
from docx import Document



def message_list(request):
    messages = Message.objects.all()  
    return render(request, 'messages/message_list.html', {'messages': messages})


def create_message(request):
    if request.method == 'POST':
        if 'file' in request.FILES:
            file = request.FILES['file']
            fs = FileSystemStorage()
            filename = fs.save(file.name, file)
            file_path = fs.url(filename)

            if file.name.endswith('.txt'):
                
                with open(file_path[1:], 'r') as f:
                    for line in f:
                        parts = line.strip().split('|')  
                        if len(parts) == 3:
                            Message.objects.create(sender=parts[0], recipient=parts[1], text=parts[2])

            elif file.name.endswith('.docx'):
                
                doc = Document(file_path[1:])
                for para in doc.paragraphs:
                    parts = para.strip().split('|')  
                    if len(parts) == 3:
                        Message.objects.create(sender=parts[0], recipient=parts[1], text=parts[2])

            elif file.name.endswith('.xlsx'):
                
                wb = openpyxl.load_workbook(file_path[1:])
                sheet = wb.active
                for row in sheet.iter_rows(min_row=2, values_only=True):  
                    if len(row) == 3:
                        sender, recipient, text = row
                        Message.objects.create(sender=sender, recipient=recipient, text=text)

            return redirect('message_list')

        else:
            form = MessageForm(request.POST)
            if form.is_valid():
                form.save()
                return redirect('message_list')

    else:
        form = MessageForm()

    return render(request, 'messages/create_message.html', {'form': form})


def test_results(request):
    
    results = TestResult.objects.all().order_by('-timestamp')  

    return render(request, 'messages/your_template.html', {'results': results})


